from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pptx import Presentation


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
}


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def is_supported_file(filename: str) -> bool:
    extension = get_file_extension(filename)
    return extension in ALLOWED_EXTENSIONS


def extract_pdf_text(file_path: str) -> str:
    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n\n".join(pages)


def extract_docx_text(file_path: str) -> str:
    document = DocxDocument(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


def extract_pptx_text(file_path: str) -> str:
    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        if slide_text:

            slides.append(
                f"Slide {slide_number}\n"
                + "\n".join(slide_text)
            )

    return "\n\n".join(slides)


def extract_txt_text(file_path: str) -> str:

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:

        return file.read()


def extract_text(
    file_path: str,
    file_extension: str,
) -> str:

    extension = file_extension.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    if extension == ".docx":
        return extract_docx_text(file_path)

    if extension == ".pptx":
        return extract_pptx_text(file_path)

    if extension == ".txt":
        return extract_txt_text(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )